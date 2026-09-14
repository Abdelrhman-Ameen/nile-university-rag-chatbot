"""Resumable public collection with a durable per-URL cache and coverage report.

Workers fetch/extract; one coordinator owns SQLite and the crawl frontier.
Only configured domains are followed. External links are inventoried for review.
"""

import heapq
import json
import sqlite3
import time
from collections import Counter
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit
from uuid import uuid4
from zipfile import ZipFile

from pypdf import PdfReader

from nu_chat.collect import (
    Collector,
    SourceBlocked,
    canonical_url,
    content_images,
    extract_html,
    priority,
)
from nu_chat.config import DATA_DIR, ROOT

HTML_EXTRACTION_VERSION = 2


def reextract_cached_html(result: dict) -> dict:
    """Apply parser fixes to saved HTML without changing download dates or refetching."""
    if result.get("html_extraction_version") == HTML_EXTRACTION_VERSION:
        return result
    if not result.get("raw_path") or not any(
        d.get("kind") == "html" for d in result.get("documents", [])
    ):
        return result
    body = Path(result["raw_path"]).read_bytes()
    url = result.get("final_url") or result["url"]
    title, text, links = extract_html(body, url)
    template = result["documents"][0]
    return {
        **result,
        "documents": [{**template, "title": title, "text": text}],
        "links": links,
        "images": content_images(body, url),
        "html_extraction_version": HTML_EXTRACTION_VERSION,
    }


def atomic_json(path: Path, value):
    temporary = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    replace_file(temporary, path)


def replace_file(temporary: Path, target: Path):
    # Windows readers/antivirus can briefly hold a file without delete sharing.
    for attempt in range(10):
        try:
            temporary.replace(target)
            return
        except PermissionError:
            if attempt == 9:
                raise
            time.sleep(0.2)


def office_text(body: bytes, kind: str) -> str:
    """Read paragraphs/table cells without executing macros or spreadsheet formulas."""
    with ZipFile(BytesIO(body)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError("Office document exceeds the expanded extraction limit")
    parts = []
    if kind == "docx":
        from docx import Document
        from docx.table import Table

        document = Document(BytesIO(body))
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
            else:
                parts.append(block.text)
    else:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(body), read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in workbook:
                parts.append("## " + sheet.title)
                for row in sheet.iter_rows():
                    cells = []
                    for cell in row:
                        value = "" if cell.value is None else str(cell.value)
                        if cell.value is not None and cell.number_format not in {
                            "General",
                            "0",
                            "@",
                        }:
                            value += f" (cell format: {cell.number_format})"
                        cells.append(value)
                    if any(cells):
                        parts.append(" | ".join(cells))
        finally:
            workbook.close()
    return "\n".join(parts)


def fetch_page(collector: Collector, url: str, raw_dir: Path) -> dict:
    fetched_at = datetime.now(timezone.utc).isoformat()
    try:
        final, body, content_type = collector.fetch(url)
        digest = sha256(body).hexdigest()
        raw_path = raw_dir / digest
        # Content-addressed raw snapshots make extraction reproducible without refetching.
        raw_path.write_bytes(body)
        is_pdf = body.startswith(b"%PDF") or "application/pdf" in content_type
        suffix = Path(urlsplit(final).path).suffix.lower()
        kind = "pdf" if is_pdf else "html"
        images, links, sections, gaps = [], [], [], []
        if is_pdf:
            reader = PdfReader(BytesIO(body))
            title = str((reader.metadata or {}).get("/Title") or Path(urlsplit(final).path).name)
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text(extraction_mode="layout") or ""
                sections.append((number, text))
                if len(text.strip()) < 100:
                    gaps.append({"page": number, "reason": "PDF page needs OCR"})
                for annotation in page.get("/Annots", []):
                    action = annotation.get_object().get("/A", {})
                    if action.get("/URI"):
                        links.append(str(action["/URI"]))
        elif "html" in content_type:
            title, text, links = extract_html(body, final)
            images = content_images(body, final)
            sections = [(None, text)]
        elif suffix in {".docx", ".xlsx"} and body.startswith(b"PK"):
            kind = suffix[1:]
            title = Path(urlsplit(final).path).name
            sections = [(None, office_text(body, kind))]
        else:
            return {
                "url": url,
                "status": "unsupported",
                "reason": content_type,
                "fetched_at": fetched_at,
                "raw_path": str(raw_path),
                "links": [],
            }
        documents = []
        for page, text in sections:
            text = text.strip()
            if len(text) < 60:
                continue
            documents.append(
                {
                    "id": sha256(f"{final}#{page}".encode()).hexdigest()[:16],
                    "url": final,
                    "title": title,
                    "text": text,
                    "page": page,
                    "kind": kind,
                    "fetched_at": fetched_at,
                    "archived": "oldwebsite." in final,
                    "content_sha256": digest,
                }
            )
        return {
            "url": url,
            "final_url": final,
            "status": "ok",
            "fetched_at": fetched_at,
            "documents": documents,
            "links": links,
            "images": images,
            "gaps": gaps,
            "raw_path": str(raw_path),
            "content_sha256": digest,
            "html_extraction_version": HTML_EXTRACTION_VERSION if kind == "html" else None,
        }
    except SourceBlocked as exc:
        return {
            "url": url,
            "status": "blocked",
            "reason": str(exc),
            "fetched_at": fetched_at,
            "links": [],
        }
    except Exception as exc:
        return {
            "url": url,
            "status": "error",
            "reason": str(exc),
            "fetched_at": fetched_at,
            "links": [],
        }


def collect(
    max_pages: int = 0,
    use_sitemaps: bool = True,
    config_path: Path = ROOT / "sources.json",
    refresh: bool = False,
    workers: int = 8,
) -> dict:
    """A zero page limit drains the discovered frontier. Cache reuse is the default.

    --refresh refetches cached URLs; failures otherwise retry after one day.
    Checkpoints preserve both the corpus and unprocessed URLs on interruption.
    """
    config = json.loads(config_path.read_text(encoding="utf-8"))
    domains = config["allowed_domains"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    cache_dir = DATA_DIR / "crawl"
    raw_dir = cache_dir / "raw"
    raw_dir.mkdir(parents=True, exist_ok=True)
    database = sqlite3.connect(cache_dir / "cache.sqlite3")
    database.execute(
        "CREATE TABLE IF NOT EXISTS pages (url TEXT PRIMARY KEY, result TEXT NOT NULL)"
    )
    database.execute(
        "CREATE TABLE IF NOT EXISTS sitemaps (origin TEXT PRIMARY KEY, result TEXT NOT NULL)"
    )
    collector = Collector(domains)
    collector.delay = float(config.get("crawl_delay_seconds", 0.4))
    documents = {}
    corpus = DATA_DIR / "documents.jsonl"
    if corpus.exists():
        for line in corpus.read_text(encoding="utf-8").splitlines():
            if line:
                doc = json.loads(line)
                if canonical_url(doc["url"], domains):
                    documents[doc["id"]] = doc
    pending, queued, seen, origins = [], set(), set(), set()
    outcomes, assets, external, gaps = {}, {}, {}, []
    requests, reused, sitemap_count = 0, 0, 0
    started_at = datetime.now(timezone.utc).isoformat()

    def enqueue(url, rank=None):
        safe = canonical_url(url, domains)
        if safe and safe not in queued:
            queued.add(safe)
            heapq.heappush(pending, (priority(safe) if rank is None else rank, len(queued), safe))

    def process(result):
        outcomes[result["url"]] = {"status": result["status"], "reason": result.get("reason")}
        if result["status"] == "ok":
            final = result["final_url"]
            # Replace this source's old extraction, preserving other sources on failure.
            stale = [
                key
                for key, doc in documents.items()
                if doc["url"] == final and doc.get("kind") != "image"
            ]
            for key in stale:
                del documents[key]
            for doc in result["documents"]:
                documents[doc["id"]] = doc
            for asset in result.get("images", []):
                assets[(final, asset["url"])] = {
                    **asset,
                    "source_url": final,
                    "source_title": result["documents"][0]["title"]
                    if result["documents"]
                    else final,
                }
            gaps.extend({"url": final, **gap} for gap in result.get("gaps", []))
        for link in result.get("links", []):
            if canonical_url(link, domains):
                enqueue(link)
            elif urlsplit(link).scheme in {"https", "http"}:
                external.setdefault(link, {"url": link, "linked_from": result["url"]})

    def checkpoint(complete=False):
        if documents:
            temporary = corpus.with_suffix(".tmp")
            unique, hashes = [], set()
            for doc in documents.values():
                digest = sha256((doc["title"] + "\n" + doc["text"]).encode()).hexdigest()
                if digest not in hashes:
                    hashes.add(digest)
                    unique.append(doc)
            temporary.write_text(
                "\n".join(json.dumps(d, ensure_ascii=False) for d in unique), encoding="utf-8"
            )
            replace_file(temporary, corpus)
        else:
            unique = []
        hosts = {}
        for url, result in outcomes.items():
            counts = hosts.setdefault(urlsplit(url).netloc, Counter())
            counts[result["status"]] += 1
        report = {
            "started_at": started_at,
            "updated_at": datetime.now(timezone.utc).isoformat(),
            "frontier_exhausted": complete,
            "network_attempts": requests,
            "cache_reused": reused,
            "attempted": len(outcomes),
            "documents": len(unique),
            "unique_urls": len({d["url"] for d in unique}),
            "sitemap_urls_discovered": sitemap_count,
            "pending_urls": len(pending) + len(inflight),
            "blocked_hosts": collector.blocked,
            "hosts": hosts,
            "images_discovered": len(assets),
            "ocr_candidates": sum(a["ocr_candidate"] for a in assets.values()),
            "pdf_pages_needing_ocr": len(gaps),
            "external_links_for_review": len(external),
            "skipped": [{"url": u, **r} for u, r in outcomes.items() if r["status"] != "ok"]
            + collector.skipped,
            "scope_note": "Public linked pages and sitemaps within configured domains. Exhausting this frontier is not proof that every university document is public or discoverable.",
        }
        atomic_json(DATA_DIR / "collection_report.json", report)
        atomic_json(
            DATA_DIR / "source_inventory.json",
            [{k: v for k, v in d.items() if k not in {"id", "text"}} for d in unique],
        )
        atomic_json(cache_dir / "assets.json", list(assets.values()))
        atomic_json(cache_dir / "pdf_ocr_queue.json", gaps)
        atomic_json(cache_dir / "external_links.json", list(external.values()))
        atomic_json(
            cache_dir / "frontier.json", [item[2] for item in pending] + list(inflight.values())
        )
        database.commit()
        return report

    for seed in config["seeds"]:
        enqueue(seed, -1)
    # Reconstruct every successful extraction from the durable cache, including
    # requests completed after the last corpus checkpoint before an interruption.
    if not refresh:
        for (payload,) in database.execute("SELECT result FROM pages"):
            cached = json.loads(payload)
            if cached["status"] == "ok" and canonical_url(cached["url"], domains):
                updated = reextract_cached_html(cached)
                if updated is not cached:
                    cached = updated
                    database.execute(
                        "UPDATE pages SET result=? WHERE url=?",
                        (json.dumps(cached, ensure_ascii=False), cached["url"]),
                    )
                process(cached)
                seen.add(cached["url"])
                reused += 1
    frontier = cache_dir / "frontier.json"
    if frontier.exists():
        for url in json.loads(frontier.read_text(encoding="utf-8")):
            enqueue(url)
    inflight = {}
    try:
        with ThreadPoolExecutor(max_workers=workers) as pool:
            while pending or inflight:
                deferred = []
                while pending and len(inflight) < workers:
                    item = heapq.heappop(pending)
                    url = item[2]
                    if url in seen:
                        continue
                    # Do not occupy every worker with requests waiting on the same host.
                    host = urlsplit(url).netloc
                    if (
                        host in collector.blocked
                        or sum(urlsplit(u).netloc == host for u in inflight.values()) >= 2
                    ):
                        deferred.append(item)
                        continue
                    origin = "https://" + urlsplit(url).netloc
                    if use_sitemaps and origin not in origins:
                        origins.add(origin)
                        row = database.execute(
                            "SELECT result FROM sitemaps WHERE origin=?", (origin,)
                        ).fetchone()
                        discovered = (
                            json.loads(row[0])
                            if row and not refresh
                            else collector.discover_sitemaps([origin])
                        )
                        database.execute(
                            "INSERT OR REPLACE INTO sitemaps VALUES (?,?)",
                            (origin, json.dumps(discovered)),
                        )
                        sitemap_count += len(discovered)
                        for link in discovered:
                            enqueue(link)
                    row = database.execute(
                        "SELECT result FROM pages WHERE url=?", (url,)
                    ).fetchone()
                    cached = json.loads(row[0]) if row and not refresh else None
                    if cached and (
                        cached["status"] == "ok"
                        or (
                            datetime.now(timezone.utc)
                            - datetime.fromisoformat(cached["fetched_at"])
                        ).total_seconds()
                        < 86400
                    ):
                        seen.add(url)
                        reused += 1
                        process(cached)
                        continue
                    if max_pages and requests >= max_pages:
                        heapq.heappush(pending, (priority(url), len(queued), url))
                        break
                    seen.add(url)
                    requests += 1
                    inflight[pool.submit(fetch_page, collector, url, raw_dir)] = url
                for item in deferred:
                    heapq.heappush(pending, item)
                if not inflight:
                    break
                completed, _ = wait(inflight, return_when=FIRST_COMPLETED)
                for future in completed:
                    url = inflight.pop(future)
                    result = future.result()
                    database.execute(
                        "INSERT OR REPLACE INTO pages VALUES (?,?)",
                        (url, json.dumps(result, ensure_ascii=False)),
                    )
                    database.commit()
                    process(result)
                    print(
                        f"[{len(outcomes)} collected; {len(pending)} queued] {result['status']}: {url}",
                        flush=True,
                    )
                if len(outcomes) % 100 < len(completed):
                    checkpoint()
        return checkpoint(complete=not pending)
    finally:
        # Save after ordinary errors/interrupts as well as normal completion.
        checkpoint(complete=not pending and not inflight)
        collector.client.close()
        database.close()
